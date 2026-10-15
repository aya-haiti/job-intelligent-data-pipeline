"""
Systeme de Recommandation a partir d'un CV (PDF ou Word).

Pipeline :
    CV.pdf/docx
        |
        v
    Extraction texte (pypdf / python-docx)
        |
        v
    Extraction competences (regex taxonomie)
        |
        v
    Construction profil candidat
        |
        v
    TF-IDF + Similarite Cosinus (reutilise RecommandeurOffres)
        |
        v
    Top N offres recommandees
"""
import os
import re
import sys
import pandas as pd

# Reutilise le moteur de recommandation deja cree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from recommandation import RecommandeurOffres
from extract_skills import SKILLS_FLAT, SKILL_TO_CATEGORY


# ============================================================================
# 1. Lecture de CV (PDF ou DOCX)
# ============================================================================
def lire_cv_pdf(chemin: str) -> str:
    """Extrait le texte d'un fichier PDF."""
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise ImportError("Installe pypdf : pip install pypdf")

    reader = PdfReader(chemin)
    texte = ""
    for page in reader.pages:
        try:
            texte += page.extract_text() + "\n"
        except Exception:
            continue
    return texte


def lire_cv_docx(chemin: str) -> str:
    """Extrait le texte d'un fichier Word."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Installe python-docx : pip install python-docx")

    doc = Document(chemin)
    texte = ""
    for paragraphe in doc.paragraphs:
        texte += paragraphe.text + "\n"
    # Tableaux dans les CV modernes
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texte += cell.text + " "
            texte += "\n"
    return texte


def lire_cv(chemin: str) -> str:
    """Lit un CV PDF ou DOCX et retourne le texte brut."""
    if not os.path.exists(chemin):
        raise FileNotFoundError(f"CV introuvable : {chemin}")

    ext = chemin.lower().split(".")[-1]
    if ext == "pdf":
        texte = lire_cv_pdf(chemin)
    elif ext in ("docx", "doc"):
        texte = lire_cv_docx(chemin)
    else:
        raise ValueError(f"Format non supporte : .{ext} (utilise PDF ou DOCX)")

    if not texte or len(texte.strip()) < 30:
        raise ValueError("Le CV semble vide ou illisible (PDF scanne ?)")

    return texte


# ============================================================================
# 2. Extraction des competences depuis le texte du CV
# ============================================================================
def extraire_competences_cv(texte_cv: str) -> list:
    """Detecte les competences techniques presentes dans le CV."""
    texte_lower = texte_cv.lower()
    competences_trouvees = []
    for skill_name, pattern in SKILLS_FLAT.items():
        try:
            if re.search(pattern, texte_lower, re.IGNORECASE):
                competences_trouvees.append(skill_name)
        except re.error:
            continue
    return competences_trouvees


def grouper_competences_par_categorie(competences: list) -> dict:
    """Regroupe les competences par categorie."""
    par_cat = {}
    for skill in competences:
        cat = SKILL_TO_CATEGORY.get(skill, "Autre")
        par_cat.setdefault(cat, []).append(skill)
    return par_cat


# ============================================================================
# 3. Pipeline complet : CV -> Recommandations
# ============================================================================
def recommander_depuis_cv(chemin_cv: str, top_n: int = 10) -> dict:
    """
    Pipeline complet : lit le CV, extrait les competences, recommande.
    Retourne un dict avec toutes les infos pour la soutenance.
    """
    print("=" * 70)
    print(f"  RECOMMANDATION A PARTIR DU CV : {os.path.basename(chemin_cv)}")
    print("=" * 70)

    # ETAPE 1 : Lecture du CV
    print("\n[1/4] Extraction du texte du CV...")
    texte_cv = lire_cv(chemin_cv)
    print(f"  {len(texte_cv)} caracteres extraits du CV")
    print(f"  Apercu (200 premiers caracteres) :")
    print(f"  {texte_cv[:200].strip()}...")

    # ETAPE 2 : Extraction des competences
    print("\n[2/4] Detection des competences techniques...")
    competences = extraire_competences_cv(texte_cv)
    print(f"  {len(competences)} competences detectees dans le CV :")

    par_categorie = grouper_competences_par_categorie(competences)
    for cat, skills in par_categorie.items():
        print(f"    [{cat}] : {', '.join(skills)}")

    if not competences:
        print("  ATTENTION : Aucune competence detectee. Verifie que ton CV contient")
        print("  des mots-cles techniques (Python, SQL, etc.)")

    # ETAPE 3 : Chargement des offres + entrainement TF-IDF
    print("\n[3/4] Chargement des offres et entrainement du modele...")
    df_offres = pd.read_csv("data/processed/all_jobs_clean.csv")
    print(f"  {len(df_offres)} offres data chargees")

    recommandeur = RecommandeurOffres()
    recommandeur.entrainer(df_offres)

    # ETAPE 4 : Recommandation
    # On utilise le texte complet du CV (richesse semantique max)
    print(f"\n[4/4] Generation des {top_n} meilleures recommandations...")
    resultats = recommandeur.recommander(texte_cv, top_n=top_n)

    return {
        "texte_cv": texte_cv,
        "competences_detectees": competences,
        "competences_par_categorie": par_categorie,
        "recommandations": resultats,
    }


# ============================================================================
# 4. Affichage et export
# ============================================================================
def afficher_resultats(resultat: dict, nom_cv: str):
    """Affiche un rapport propre dans la console."""
    print("\n" + "=" * 70)
    print("  RESUME DU PROFIL EXTRAIT DU CV")
    print("=" * 70)
    print(f"  Competences techniques detectees : {len(resultat['competences_detectees'])}")
    print(f"  Categories couvertes              : {len(resultat['competences_par_categorie'])}")

    print("\n" + "=" * 70)
    print(f"  TOP {len(resultat['recommandations'])} OFFRES RECOMMANDEES")
    print("=" * 70)
    print(f"  {'Rang':<5} {'Score':>6}  {'Titre':<38} {'Source':<10} {'Ville'}")
    print(f"  {'-'*5} {'-'*6}  {'-'*38} {'-'*10} {'-'*20}")
    for _, row in resultat["recommandations"].iterrows():
        titre = str(row["titre"])[:37]
        ville = str(row["ville"])[:20]
        print(f"  {int(row['rang']):<5} {row['score_pct']:>5.1f}%  {titre:<38} {str(row['source']):<10} {ville}")

    # Export CSV
    os.makedirs("data/processed", exist_ok=True)
    export_path = f"data/processed/recommandations_cv_{nom_cv}.csv"

    df_export = resultat["recommandations"].copy()
    df_export["cv_source"] = nom_cv
    df_export["competences_cv"] = " | ".join(resultat["competences_detectees"])
    df_export.to_csv(export_path, index=False, encoding="utf-8-sig")

    print(f"\n  Export Power BI : {export_path}")
    print("=" * 70)


# ============================================================================
# Main
# ============================================================================
if __name__ == "__main__":
    # CV par defaut, modifiable en argument
    chemin_cv = sys.argv[1] if len(sys.argv) > 1 else "data/cv/mon_cv.pdf"

    nom_cv = os.path.splitext(os.path.basename(chemin_cv))[0]

    try:
        resultat = recommander_depuis_cv(chemin_cv, top_n=10)
        afficher_resultats(resultat, nom_cv)
        print("\n  Pipeline CV -> Recommandation termine avec succes !\n")
    except FileNotFoundError as e:
        print(f"\n  ERREUR : {e}")
        print("  Place ton CV dans : data/cv/mon_cv.pdf")
        print("  Ou lance avec : python scripts/nlp/recommandation_cv.py chemin/vers/cv.pdf")
    except Exception as e:
        print(f"\n  ERREUR : {e}")
        import traceback
        traceback.print_exc()